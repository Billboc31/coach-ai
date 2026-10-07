"""Excel preview, explicit column mapping and persistent workout set logs."""

import base64
import hashlib
import io
import json
import re
import threading
import unicodedata
import uuid
import zipfile
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import text

from coach import db
from coach.config import user_id

lock = threading.Lock()
MAX_FILE = 5 * 1024 * 1024


def normalized(value):
    return "".join(
        c for c in unicodedata.normalize("NFKD", str(value).lower()) if not unicodedata.combining(c)
    ).strip()


def video_id(url):
    if not url:
        return None
    p = urlparse(url)
    if p.scheme != "https":
        raise ValueError("Lien YouTube HTTPS attendu.")
    host = p.hostname or ""
    if host == "youtu.be":
        key = p.path.strip("/").split("/")[0]
    elif host in {"youtube.com", "www.youtube.com", "m.youtube.com"}:
        key = parse_qs(p.query).get("v", [""])[0]
        if p.path.startswith(("/embed/", "/shorts/")):
            key = p.path.split("/")[2]
    else:
        raise ValueError("Lien YouTube attendu.")
    if not re.fullmatch(r"[A-Za-z0-9_-]{11}", key):
        raise ValueError("Lien vidéo YouTube invalide.")
    return key


ALIASES = {
    "name": ["exercice", "exercise", "exercices", "nom"],
    "sets": ["series", "sets", "serie"],
    "reps": ["repetitions", "reps", "rep"],
    "weight": ["poids", "charge", "kg", "poids kg", "charge kg", "poids e"],
    "rest": ["repos", "rest", "recuperation"],
    "notes": ["consignes", "notes", "tempo"],
    "video": ["video", "youtube", "lien video"],
}


def preview(filename, encoded):
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Choisir un fichier .xlsx ; convertir les anciens .xls dans Excel.")
    try:
        binary = base64.b64decode(encoded, validate=True)
        if len(binary) > MAX_FILE:
            raise ValueError()
        with zipfile.ZipFile(io.BytesIO(binary)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
                raise ValueError()
        workbook = load_workbook(
            io.BytesIO(binary), read_only=False, data_only=False, keep_links=False
        )
    except Exception:
        raise ValueError("Fichier Excel invalide ou trop volumineux (5 Mo maximum).") from None
    sheets = []
    try:
        if len(workbook.worksheets) > 30:
            raise ValueError("30 feuilles maximum.")
        for sheet in workbook.worksheets[:30]:
            rows = []
            if sheet.max_row > 1000 or sheet.max_column > 40:
                raise ValueError("1000 lignes et 40 colonnes maximum par feuille.")
            links, dates = {}, []
            for i, cells in enumerate(sheet.iter_rows()):
                row = [c.value for c in cells]
                for c in cells:
                    if c.hyperlink:
                        links[f"{c.row}:{c.column - 1}"] = c.hyperlink.target
                    if getattr(c, "is_date", False):
                        dates.append(f"{c.row}:{c.column - 1}")
                if i == 1000:
                    raise ValueError(
                        "Une feuille dépasse 1000 lignes. Réduis le fichier avant import."
                    )
                rows.append([str(v) if v is not None else "" for v in row])
            while rows and not any(rows[-1]):
                rows.pop()
            expanded = [list(row) for row in rows]
            for merged in sheet.merged_cells.ranges:
                if merged.min_col == merged.max_col and merged.min_row <= len(rows):
                    col = merged.min_col - 1
                    for j in range(merged.min_row, min(merged.max_row, len(rows))):
                        expanded[j][col] = rows[merged.min_row - 1][col]
            best = 0
            mapping = {}
            for i, row in enumerate(rows[:30]):
                candidate = {
                    field: next(
                        (col for col, value in enumerate(row) if normalized(value) in names), None
                    )
                    for field, names in ALIASES.items()
                }
                if candidate.get("name") is not None and sum(
                    v is not None for v in candidate.values()
                ) > len(mapping):
                    best, mapping = i, {k: v for k, v in candidate.items() if v is not None}
            if mapping:
                for col, value in enumerate(rows[best]):
                    label = normalized(value)
                    if label.startswith("consigne"):
                        mapping["notes"] = col
                    if value.strip() == "R":
                        mapping["rest"] = col
                    if value.strip() == "r":
                        mapping["short_rest"] = col
                    if label.startswith("lien video"):
                        mapping["video"] = col
                if mapping.get("order") == 2 and any(
                    normalized(row[1]) in {"echauf.", "echauf", "renfo"}
                    for row in rows[best + 1 : best + 30]
                ):
                    mapping["group"] = 1
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": rows,
                    "expanded": expanded,
                    "links": links,
                    "date_cells": dates,
                    "header": best,
                    "mapping": mapping,
                }
            )
    finally:
        workbook.close()
    value = {
        "id": uuid.uuid4().hex,
        "filename": filename[:150],
        "sheets": sheets,
        "created_at": db.now(),
        "fingerprint": hashlib.sha256(binary).hexdigest(),
    }
    with db.connection() as conn:
        existing = conn.execute(
            text(
                "SELECT data FROM records WHERE user_id=:u AND kind='gym_import' "
                "AND json_extract(data,'$.fingerprint')=:f LIMIT 1"
            ),
            {"u": user_id(), "f": value["fingerprint"]},
        ).first()
    if existing:
        return json.loads(existing[0])
    db.upsert_record("gym_source", value["id"], {"filename": filename, "base64": encoded})
    db.upsert_record("gym_import", value["id"], value)
    return value


class SheetMapping(BaseModel):
    name: str = Field(max_length=100)
    header: int = Field(ge=0, le=999)
    mapping: dict[str, int] = Field(max_length=16)


class ImportSelection(BaseModel):
    title: str = Field(min_length=1, max_length=120)
    performances_are_kg: bool = False
    sheets: list[SheetMapping] = Field(min_length=1, max_length=30)


def numeric(value, *, kg=False):
    if kg:
        value = re.sub(r"\s*kg\s*$", "", str(value), flags=re.I)
    try:
        parsed = float(str(value).replace(",", "."))
        return parsed if 0 <= parsed <= 2000 else None
    except ValueError:
        return None


def confirm(key, selection):
    with lock:
        source = db.record("gym_import", key)
        if not source:
            return None
        if source.get("program_id"):
            return db.record("gym_program", source["program_id"])
        days = []
        imported = []
        exercises = {}
        for chosen in selection.sheets:
            sheet = next((s for s in source["sheets"] if s["name"] == chosen.name), None)
            if not sheet or chosen.header >= len(sheet["rows"]) or "name" not in chosen.mapping:
                raise ValueError(
                    "Choisir la ligne des titres et la colonne exercice pour chaque feuille."
                )
            if any(k not in ALIASES or not 0 <= v < 40 for k, v in chosen.mapping.items()):
                raise ValueError("Association de colonnes invalide.")
            items = []
            for index, row, label, week, session in workout_rows(sheet, chosen):

                def cell(field):
                    col = chosen.mapping.get(field)
                    return row[col].strip() if col is not None and col < len(row) else ""

                name = cell("name")
                if not name or name.startswith("="):
                    continue
                eid = hashlib.sha256(normalized(name).encode()).hexdigest()[:24]
                count = numeric(cell("sets"))
                count = (
                    int(count)
                    if count is not None and count.is_integer() and 1 <= count <= 20
                    else None
                )
                media = sheet.get("links", {}).get(
                    f"{index}:{chosen.mapping.get('video')}", cell("video")
                )
                try:
                    vid = video_id(media)
                except ValueError:
                    vid = None
                exercises[eid] = {"id": eid, "name": name[:200], "video_id": vid}
                perf_col = chosen.mapping.get("performance")
                perf = []
                if perf_col is not None:
                    end = next(
                        (
                            j
                            for j in range(perf_col + 1, len(row))
                            if sheet["rows"][chosen.header][j]
                        ),
                        len(row),
                    )
                    perf = [v for v in row[perf_col:end] if v]
                if perf:
                    imported.append(
                        {
                            "id": hashlib.sha256(
                                json.dumps(
                                    [
                                        source["filename"],
                                        chosen.name,
                                        index,
                                        eid,
                                        perf,
                                        cell("reps"),
                                    ],
                                    ensure_ascii=False,
                                ).encode()
                            ).hexdigest(),
                            "exercise_id": eid,
                            "source": "excel",
                            "sheet": chosen.name,
                            "week": week,
                            "session": session,
                            "row": index,
                            "performance": perf,
                            "reps": cell("reps"),
                            "weights_kg": [numeric(v, kg=True) for v in perf]
                            if selection.performances_are_kg
                            else None,
                            "imported_at": db.now(),
                        }
                    )
                items.append(
                    {
                        "id": uuid.uuid4().hex,
                        "exercise_id": eid,
                        "name": name[:200],
                        "sets": count,
                        "reps": cell("reps"),
                        "weight": numeric(cell("weight"), kg=True),
                        "rest": cell("rest"),
                        "notes": cell("notes"),
                        "comments": cell("comments"),
                        "rm": cell("rm"),
                        "short_rest": cell("short_rest"),
                        "group": cell("group"),
                        "label": label,
                        "sheet": chosen.name,
                        "week": week,
                        "order": cell("order"),
                        "tempo": cell("tempo"),
                        "performance": perf,
                        "media_url": safe_media(media),
                        "warnings": [
                            f"{field}: cellule Excel de type date, à vérifier."
                            for field in ("reps", "sets")
                            if f"{index}:{chosen.mapping.get(field)}" in sheet.get("date_cells", [])
                        ],
                        "source_row": index,
                        "source_values": sheet["rows"][index - 1],
                    }
                )
            grouped = {}
            for item in items:
                grouped.setdefault(item["label"], []).append(item)
            for label, group in grouped.items():
                days.append(
                    {
                        "id": uuid.uuid4().hex,
                        "name": label,
                        "sheet": chosen.name,
                        "week": group[0]["week"],
                        "exercises": group,
                    }
                )
        if not days:
            raise ValueError("Aucun exercice identifié. Vérifie les colonnes avant de confirmer.")
        program = {
            "id": key,
            "title": selection.title.strip(),
            "days": days,
            "created_at": db.now(),
            "source": source["filename"],
            "source_sheets": source["sheets"],
        }
        # Exercise history never gets replaced by re-importing a program.
        for eid, exercise in exercises.items():
            old = db.record("gym_exercise", eid)
            db.upsert_record(
                "gym_exercise",
                eid,
                {**exercise, "video_id": old.get("video_id") or exercise["video_id"]},
            )
        for entry in imported:
            db.upsert_record("gym_excel_history", entry["id"], entry)
        db.upsert_record("gym_program", key, program)
        db.upsert_record("gym_import", key, {**source, "program_id": key})
        return program


def history(exercise_id):
    with db.connection() as conn:
        rows = conn.execute(
            text(
                "SELECT w.data,e.value FROM records w, json_each(w.data,'$.exercises') e "
                "WHERE w.user_id=:u AND w.kind='gym_workout' AND json_extract(e.value,'$.exercise_id')=:id "
                "AND EXISTS (SELECT 1 FROM json_each(e.value,'$.logged_sets') s "
                "WHERE json_extract(s.value,'$.done')=1 AND (json_extract(s.value,'$.reps') IS NOT NULL "
                "OR json_extract(s.value,'$.seconds') IS NOT NULL)) "
                "ORDER BY json_extract(w.data,'$.started_at') DESC LIMIT 20"
            ),
            {"u": user_id(), "id": exercise_id},
        ).all()
    result = []
    for raw, exercise_raw in rows:
        workout, exercise = json.loads(raw), json.loads(exercise_raw)
        performed = [
            s
            for s in exercise["logged_sets"]
            if s["done"] and (s["reps"] is not None or s.get("seconds") is not None)
        ]
        result.append(
            {
                "workout_id": workout["id"],
                "date": workout["started_at"],
                "sets": performed,
                "finished": bool(workout.get("finished_at")),
            }
        )
    return result


def recent_workouts(limit=50):
    with db.connection() as conn:
        rows = conn.execute(
            text(
                "SELECT data FROM records WHERE user_id=:u AND kind='gym_workout' "
                "ORDER BY json_extract(data,'$.started_at') DESC LIMIT :l"
            ),
            {"u": user_id(), "l": limit},
        ).all()
    return [json.loads(r[0]) for r in rows]


def overview():
    return {
        "programs": [
            {k: v for k, v in r["data"].items() if k != "source_sheets"}
            for r in db.records("gym_program", 30)
        ],
        "exercises": [r["data"] for r in db.records("gym_exercise", 1000)],
        "workouts": recent_workouts(),
    }


def start(program_id, day_id):
    with lock:
        program = db.record("gym_program", program_id)
        day = next((d for d in program.get("days", []) if d["id"] == day_id), None)
        if not day:
            raise ValueError("Programme ou séance introuvable.")
        rows = []
        for item in day["exercises"]:
            past = history(item["exercise_id"])
            reference = past[0] if past else None
            logs = [
                {"weight": None, "reps": None, "seconds": None, "done": False}
                for _ in range(item["sets"] or 1)
            ]
            rows.append(
                {
                    **item,
                    "logged_sets": logs,
                    "reference": reference,
                    "excel_history": imported_history(item["exercise_id"]),
                }
            )
        value = {
            "id": uuid.uuid4().hex,
            "program_id": program_id,
            "title": day["name"],
            "started_at": db.now(),
            "finished_at": None,
            "revision": 1,
            "exercises": rows,
        }
        db.upsert_record("gym_workout", value["id"], value)
        return value


class LoggedSet(BaseModel):
    weight: float | None = Field(default=None, ge=0, le=2000, allow_inf_nan=False)
    reps: int | None = Field(default=None, ge=0, le=1000)
    seconds: int | None = Field(default=None, ge=0, le=7200)
    done: bool = False


class WorkoutUpdate(BaseModel):
    revision: int = Field(ge=1)
    sets: dict[str, list[LoggedSet]] = Field(max_length=1000)
    finish: bool = False


class Conflict(ValueError):
    pass


def update_workout(key, body):
    with lock:
        value = db.record("gym_workout", key)
        if not value:
            return None
        if value["revision"] != body.revision:
            raise Conflict("Séance modifiée ailleurs. Recharge avant de continuer.")
        if value.get("finished_at"):
            raise Conflict("Séance terminée ; son historique est conservé.")
        if set(body.sets) != {e["id"] for e in value["exercises"]}:
            raise ValueError("Liste des exercices incohérente.")
        for exercise in value["exercises"]:
            sets = body.sets[exercise["id"]]
            if not 1 <= len(sets) <= 30 or any(
                s.done and s.reps is None and s.seconds is None for s in sets
            ):
                raise ValueError(
                    "Renseigne les répétitions des séries validées (30 séries maximum)."
                )
            exercise["logged_sets"] = [s.model_dump() for s in sets]
        value.update(revision=value["revision"] + 1, updated_at=db.now())
        if body.finish:
            if not any(s["done"] for e in value["exercises"] for s in e["logged_sets"]):
                raise ValueError("Valide au moins une série avant de terminer.")
            value["finished_at"] = db.now()
        db.upsert_record("gym_workout", key, value)
        return value


def set_video(key, url):
    with lock:
        value = db.record("gym_exercise", key)
        if not value:
            return None
        value["video_id"] = video_id(url)
        db.upsert_record("gym_exercise", key, value)
        return value


ALIASES.update(
    {
        "comments": ["commentaires de seance", "commentaires"],
        "rm": ["rm"],
        "short_rest": ["repos court"],
        "group": ["groupe", "phase"],
        "order": ["ordre", "order"],
        "tempo": ["tempo"],
        "performance": ["performance", "performances"],
    }
)


def safe_media(url):
    p = urlparse(url or "")
    if p.username or p.password:
        return None
    return (
        url
        if p.scheme == "https"
        and p.hostname
        in {"1drv.ms", "onedrive.live.com", "youtu.be", "youtube.com", "www.youtube.com"}
        else None
    )


def workout_rows(sheet, chosen):
    week, session = "", chosen.name
    structured = "order" in chosen.mapping
    for index, row in enumerate(
        sheet.get("expanded", sheet["rows"])[chosen.header + 1 :], chosen.header + 2
    ):
        name = row[chosen.mapping["name"]].strip()
        marker = next((v for v in row if re.fullmatch(r"semaine\s+\d+", normalized(v))), None)
        if marker and not name:
            week = marker
            continue
        if re.match(r"^seance\s+\d+", normalized(name)):
            session = name
            continue
        if (
            normalized(name).startswith("consignes generales")
            or normalized(name) in ALIASES["name"]
        ):
            continue
        if structured and not row[chosen.mapping["order"]].strip():
            continue
        label = " · ".join(
            v for v in (chosen.name, week, session if session != chosen.name else "") if v
        )
        yield index, row, label, week, session


def imported_history(exercise_id):
    with db.connection() as conn:
        rows = conn.execute(
            text(
                "SELECT data FROM records WHERE user_id=:u AND kind='gym_excel_history' "
                "AND json_extract(data,'$.exercise_id')=:e ORDER BY updated_at DESC, "
                "json_extract(data,'$.row') DESC LIMIT 30"
            ),
            {"u": user_id(), "e": exercise_id},
        ).all()
    return [json.loads(r[0]) for r in rows]


def context():
    return [
        {
            "date": w["started_at"],
            "completed": bool(w["finished_at"]),
            "title": w["title"],
            "units": {"weight": "kg", "reps": "count", "seconds": "seconds"},
            "exercises": [
                {"name": e["name"], "sets": [s for s in e["logged_sets"] if s["done"]]}
                for e in w["exercises"][:25]
            ],
        }
        for w in recent_workouts(5)
    ]
